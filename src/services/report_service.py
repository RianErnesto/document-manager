"""
Serviço para geração de relatórios em PDF e XLSX.
"""
import os
from datetime import datetime
from typing import List, Optional, Tuple

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from ..models.document import Document
from ..config import COMPANY_NAME


class ReportService:
    """Serviço para geração de relatórios."""

    @staticmethod
    def generate_pdf(
        documents: List[Document],
        file_path: str,
        date_start: Optional[datetime] = None,
        date_end: Optional[datetime] = None,
    ) -> Tuple[bool, str]:
        """
        Gera relatório em PDF.

        Args:
            documents: Lista de documentos
            file_path: Caminho do arquivo
            date_start: Data inicial do filtro
            date_end: Data final do filtro

        Returns:
            (sucesso, mensagem)
        """
        try:
            # Configuração do documento PDF
            pdf_doc = SimpleDocTemplate(
                file_path,
                pagesize=landscape(A4),
                rightMargin=1.5 * cm,
                leftMargin=1.5 * cm,
                topMargin=1.5 * cm,
                bottomMargin=1.5 * cm,
            )

            elements = []
            styles = getSampleStyleSheet()

            # Título
            title_style = ParagraphStyle(
                "CustomTitle",
                parent=styles["Heading1"],
                fontSize=18,
                spaceAfter=20,
                alignment=1,  # Center
            )
            elements.append(Paragraph(f"{COMPANY_NAME} - Relatório de Documentos", title_style))

            # Período
            period_text = ReportService._get_period_text(date_start, date_end)
            period_style = ParagraphStyle(
                "Period",
                parent=styles["Normal"],
                fontSize=10,
                spaceAfter=20,
                alignment=1,
            )
            elements.append(Paragraph(period_text, period_style))

            # Data de geração
            gen_date = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            elements.append(Paragraph(gen_date, period_style))
            elements.append(Spacer(1, 20))

            if not documents:
                elements.append(Paragraph("Nenhum documento encontrado no período.", styles["Normal"]))
            else:
                # Tabela
                data = [["ID", "Código QR", "Estante", "Prateleira", "Caixa", "Classificação", "Data Cadastro"]]

                for doc in documents:
                    data.append([
                        str(doc.id),
                        doc.qr_code,
                        str(doc.rack),
                        str(doc.shelf),
                        str(doc.box),
                        doc.classification,
                        doc.created_at.strftime("%d/%m/%Y %H:%M") if doc.created_at else "",
                    ])

                # Estilo da tabela
                table = Table(data, colWidths=[40, 150, 60, 70, 60, 90, 110])
                table.setStyle(TableStyle([
                    # Header
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a73e8")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 11),
                    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("TOPPADDING", (0, 0), (-1, 0), 12),

                    # Body
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 10),
                    ("ALIGN", (0, 1), (-1, -1), "CENTER"),
                    ("BOTTOMPADDING", (0, 1), (-1, -1), 8),
                    ("TOPPADDING", (0, 1), (-1, -1), 8),

                    # Borders
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#dadce0")),

                    # Alternate row colors
                    *[
                        ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f8f9fa"))
                        for i in range(2, len(data), 2)
                    ],
                ]))

                elements.append(table)
                elements.append(Spacer(1, 20))

                # Total
                total_text = f"Total de documentos: {len(documents)}"
                total_style = ParagraphStyle(
                    "Total",
                    parent=styles["Normal"],
                    fontSize=11,
                    fontName="Helvetica-Bold",
                )
                elements.append(Paragraph(total_text, total_style))

            # Gera o PDF
            pdf_doc.build(elements)
            return True, f"Relatório PDF gerado com sucesso!\n{file_path}"

        except Exception as e:
            return False, f"Erro ao gerar PDF: {str(e)}"

    @staticmethod
    def generate_xlsx(
        documents: List[Document],
        file_path: str,
        date_start: Optional[datetime] = None,
        date_end: Optional[datetime] = None,
    ) -> Tuple[bool, str]:
        """
        Gera relatório em XLSX.

        Args:
            documents: Lista de documentos
            file_path: Caminho do arquivo
            date_start: Data inicial do filtro
            date_end: Data final do filtro

        Returns:
            (sucesso, mensagem)
        """
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Documentos"

            # Estilos
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="1a73e8", end_color="1a73e8", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")

            cell_alignment = Alignment(horizontal="center", vertical="center")
            cell_border = Border(
                left=Side(style="thin", color="dadce0"),
                right=Side(style="thin", color="dadce0"),
                top=Side(style="thin", color="dadce0"),
                bottom=Side(style="thin", color="dadce0"),
            )

            alt_fill = PatternFill(start_color="f8f9fa", end_color="f8f9fa", fill_type="solid")

            # Título
            ws.merge_cells("A1:G1")
            ws["A1"] = f"{COMPANY_NAME} - Relatório de Documentos"
            ws["A1"].font = Font(bold=True, size=16)
            ws["A1"].alignment = Alignment(horizontal="center")

            # Período
            ws.merge_cells("A2:G2")
            ws["A2"] = ReportService._get_period_text(date_start, date_end)
            ws["A2"].alignment = Alignment(horizontal="center")

            # Data de geração
            ws.merge_cells("A3:G3")
            ws["A3"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            ws["A3"].alignment = Alignment(horizontal="center")

            # Espaço
            start_row = 5

            # Headers
            headers = ["ID", "Código QR", "Estante", "Prateleira", "Caixa", "Classificação", "Data Cadastro"]
            col_widths = [10, 30, 12, 14, 12, 18, 20]

            for col, (header, width) in enumerate(zip(headers, col_widths), 1):
                cell = ws.cell(row=start_row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = cell_border
                ws.column_dimensions[get_column_letter(col)].width = width

            # Dados
            for row_idx, doc in enumerate(documents, start_row + 1):
                row_data = [
                    doc.id,
                    doc.qr_code,
                    doc.rack,
                    doc.shelf,
                    doc.box,
                    doc.classification,
                    doc.created_at.strftime("%d/%m/%Y %H:%M") if doc.created_at else "",
                ]

                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    cell.alignment = cell_alignment
                    cell.border = cell_border

                    # Cor alternada
                    if (row_idx - start_row) % 2 == 0:
                        cell.fill = alt_fill

            # Total
            total_row = start_row + len(documents) + 2
            ws.cell(row=total_row, column=1, value=f"Total de documentos: {len(documents)}")
            ws.cell(row=total_row, column=1).font = Font(bold=True)

            # Salva
            wb.save(file_path)
            return True, f"Relatório XLSX gerado com sucesso!\n{file_path}"

        except Exception as e:
            return False, f"Erro ao gerar XLSX: {str(e)}"

    @staticmethod
    def _get_period_text(
        date_start: Optional[datetime],
        date_end: Optional[datetime],
    ) -> str:
        """Retorna texto do período."""
        if date_start and date_end:
            return f"Período: {date_start.strftime('%d/%m/%Y')} a {date_end.strftime('%d/%m/%Y')}"
        elif date_start:
            return f"Período: a partir de {date_start.strftime('%d/%m/%Y')}"
        elif date_end:
            return f"Período: até {date_end.strftime('%d/%m/%Y')}"
        else:
            return "Período: Todos os documentos"
